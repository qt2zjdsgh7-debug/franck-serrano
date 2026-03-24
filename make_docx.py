#!/usr/bin/env python3
"""Convert markdown files to styled DOCX using python-docx."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Colors ──────────────────────────────────────────────────────────────────
DARK_BLUE   = RGBColor(0x0F, 0x34, 0x60)
MID_BLUE    = RGBColor(0x16, 0x21, 0x3E)
RED_ACCENT  = RGBColor(0xE9, 0x45, 0x60)
PURPLE      = RGBColor(0x53, 0x34, 0x83)
LIGHT_TEXT  = RGBColor(0x2D, 0x2D, 0x44)
GRAY        = RGBColor(0x6B, 0x72, 0x80)
CODE_BG     = RGBColor(0x1E, 0x1E, 0x2E)
CODE_FG     = RGBColor(0xCD, 0xD6, 0xF4)
TABLE_HDR   = RGBColor(0x0F, 0x34, 0x60)
TABLE_ALT   = RGBColor(0xF7, 0xF9, 0xFF)
QUOTE_BG    = RGBColor(0xF0, 0xF4, 0xFF)


def set_cell_bg(cell, color: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    tcPr.append(shd)


def add_border_left(para, color: RGBColor, size: int = 24):
    """Add a left border to a paragraph (for blockquotes / h2)."""
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), str(size))
    left.set(qn('w:space'), '10')
    left.set(qn('w:color'), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    pBdr.append(left)
    pPr.append(pBdr)


def set_para_bg(para, color: RGBColor):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    pPr.append(shd)


def set_margins(doc, top=2.5, bottom=2.5, left=3.0, right=3.0):
    for section in doc.sections:
        section.top_margin    = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin   = Cm(left)
        section.right_margin  = Cm(right)


def style_normal(doc):
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.font.color.rgb = LIGHT_TEXT


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), f"{DARK_BLUE[0]:02X}{DARK_BLUE[1]:02X}{DARK_BLUE[2]:02X}")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.space_after = Pt(16)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = MID_BLUE
    add_border_left(p, RED_ACCENT, size=32)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after  = Pt(8)
    return p


def add_h3(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after  = Pt(6)
    return p


def add_h4(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = PURPLE
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    return p


def add_blockquote(doc, text):
    p = doc.add_paragraph()
    # Strip leading >
    text = re.sub(r'^>\s*', '', text).strip()
    _add_inline(p, text, base_color=RGBColor(0x3A, 0x3A, 0x5C))
    p.paragraph_format.left_indent  = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(6)
    set_para_bg(p, QUOTE_BG)
    add_border_left(p, DARK_BLUE, size=24)
    return p


def add_code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = CODE_FG
    p.paragraph_format.left_indent  = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    set_para_bg(p, CODE_BG)
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f"{0xE8:02X}{0xEA:02X}{0xF6:02X}")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(12)


def _add_inline(para, text, base_color=LIGHT_TEXT):
    """Parse inline **bold**, `code`, and plain text within a paragraph."""
    pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')
    parts = pattern.split(text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.font.bold = True
            run.font.color.rgb = DARK_BLUE
        elif part.startswith('`') and part.endswith('`'):
            run = para.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            run = para.add_run(part)
            run.font.color.rgb = base_color
    for run in para.runs:
        if not run.font.name or run.font.name == 'Calibri':
            run.font.name = 'Calibri'


def add_normal(doc, text):
    p = doc.add_paragraph()
    _add_inline(p, text)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    _add_inline(p, text)
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_table(doc, header_row, rows):
    n_cols = len(header_row)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.style = 'Table Grid'
    # Header
    hdr = table.rows[0]
    for i, cell_text in enumerate(header_row):
        cell = hdr.cells[i]
        set_cell_bg(cell, TABLE_HDR)
        run = cell.paragraphs[0].add_run(cell_text.strip())
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
    # Body
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        bg = TABLE_ALT if r_idx % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            _add_inline(p, cell_text.strip())
            for run in p.runs:
                run.font.size = Pt(10)
    doc.add_paragraph()  # spacing after table


# ── Main parser ──────────────────────────────────────────────────────────────

def parse_table_line(line):
    return [c for c in line.strip().strip('|').split('|')]


def is_separator_row(row):
    return all(re.match(r'^[-: ]+$', c.strip()) for c in row if c.strip())


def convert(md_path: Path, docx_path: Path):
    doc = Document()
    set_margins(doc)
    style_normal(doc)

    lines = md_path.read_text(encoding='utf-8').splitlines()
    i = 0
    in_code = False
    code_lines = []
    table_rows = []
    table_header = None

    while i < len(lines):
        line = lines[i]

        # ── YAML frontmatter (skip) ──────────────────────────────────────
        if i == 0 and line.strip() == '---':
            i += 1
            while i < len(lines) and lines[i].strip() != '---':
                i += 1
            i += 1
            continue

        # ── Code fences ─────────────────────────────────────────────────
        if line.startswith('```') or line.startswith('````'):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, '\n'.join(code_lines))
                in_code = False
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # ── Flush pending table ─────────────────────────────────────────
        def flush_table():
            nonlocal table_header, table_rows
            if table_header:
                add_table(doc, table_header, table_rows)
            table_header = None
            table_rows = []

        # ── Table rows ──────────────────────────────────────────────────
        if '|' in line and line.strip().startswith('|'):
            cells = parse_table_line(line)
            if table_header is None:
                table_header = cells
            elif is_separator_row(cells):
                pass  # skip separator
            else:
                table_rows.append(cells)
            i += 1
            continue
        else:
            flush_table()

        # ── HR ───────────────────────────────────────────────────────────
        if re.match(r'^---+$', line.strip()) or re.match(r'^===+$', line.strip()):
            add_hr(doc)
            i += 1
            continue

        # ── Headings ────────────────────────────────────────────────────
        if line.startswith('#### '):
            add_h4(doc, line[5:].strip())
        elif line.startswith('### '):
            add_h3(doc, line[4:].strip())
        elif line.startswith('## '):
            add_h2(doc, line[3:].strip())
        elif line.startswith('# '):
            add_title(doc, line[2:].strip())

        # ── Blockquotes ─────────────────────────────────────────────────
        elif line.startswith('>'):
            add_blockquote(doc, line)

        # ── List items ──────────────────────────────────────────────────
        elif re.match(r'^(\s*)[-*+] ', line):
            m = re.match(r'^(\s*)[-*+] (.*)', line)
            level = len(m.group(1)) // 2
            add_bullet(doc, m.group(2).strip(), level=level)

        elif re.match(r'^\d+\. ', line):
            add_bullet(doc, re.sub(r'^\d+\. ', '', line).strip())

        # ── Blank line ───────────────────────────────────────────────────
        elif line.strip() == '':
            pass  # natural spacing

        # ── Normal paragraph ─────────────────────────────────────────────
        else:
            add_normal(doc, line.strip())

        i += 1

    flush_table()  # flush any trailing table
    doc.save(str(docx_path))
    print(f"  Saved → {docx_path}  ({docx_path.stat().st_size // 1024} KB)")


FILES = [
    {
        "md":   Path("/home/user/franck-serrano/expert-prompt-guide.md"),
        "docx": Path("/home/user/franck-serrano/documents/expert-prompt-guide.docx"),
    },
    {
        "md":   Path("/home/user/franck-serrano/.claude/skills/expert-prompt/SKILL.md"),
        "docx": Path("/home/user/franck-serrano/documents/expert-prompt-skill.docx"),
    },
]

Path("/home/user/franck-serrano/documents").mkdir(exist_ok=True)

for f in FILES:
    print(f"Converting: {f['md'].name}")
    convert(f['md'], f['docx'])

print("\nDone.")
